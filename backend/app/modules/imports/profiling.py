"""Full-volume source audit. Detailed results stay in the ignored cache directory."""
from collections import Counter, defaultdict
from dataclasses import asdict
from datetime import date, datetime, timezone
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import time

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from .adapters.common import ADAPTER_VERSION, iter_workbook, file_hash


def json_default(value):
    if isinstance(value, ArrayFormula):
        return {"formula": value.text, "ref": value.ref}
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(type(value).__name__)


def dumps(value):
    return json.dumps(value, ensure_ascii=False, default=json_default, sort_keys=True)


def profile_sources(root: Path, output: Path) -> dict:
    started = time.perf_counter()
    output.mkdir(parents=True, exist_ok=True)
    report = {"adapter_version": ADAPTER_VERSION, "started_at": datetime.now(timezone.utc).isoformat(),
              "files": [], "suppliers": {}}
    for folder, supplier in (("IEK", "iek"), ("Systeme electric", "systeme")):
        sets = defaultdict(set)
        series = defaultdict(dict)
        transactions = defaultdict(lambda: Decimal(0))
        transaction_units = defaultdict(set)
        constraints, sales_constraints = {}, {}
        supplier_stats = Counter()
        for path in sorted((root / folder).glob("*.xlsx")):
            digest = file_hash(path)
            item = {"file": path.relative_to(root).as_posix(), "sha256": digest, "sheets": [],
                    "kinds": {}, "issues": [], "normalization_errors": 0}
            # Content-addressed staging output is diagnostic, not an activated dataset.
            target = output / f"{supplier}-{digest}-v{ADAPTER_VERSION}.jsonl"
            temporary = target.with_suffix(".tmp")
            counters = defaultdict(Counter)
            codes = defaultdict(set)
            seen = defaultdict(set)
            first_date, last_date = None, None
            with temporary.open("w", encoding="utf-8") as stream:
                for row in iter_workbook(path, supplier):
                    stream.write(dumps(asdict(row)) + "\n")
                    count = counters[row.kind]
                    count["rows"] += 1
                    if row.errors:
                        count["errors"] += 1
                        item["normalization_errors"] += 1
                        if len(item["issues"]) < 20:
                            item["issues"].append({"sheet": row.sheet, "row": row.row, "errors": row.errors})
                    if row.code is None or row.errors:
                        continue
                    fields = row.fields
                    row_key = hashlib.sha256(dumps(fields).encode()).hexdigest()
                    if row_key in seen[row.kind]:
                        count["identical_rows"] += 1
                    seen[row.kind].add(row_key)
                    if row.code in codes[row.kind]:
                        count["repeated_code_rows"] += 1
                    codes[row.kind].add(row.code)
                    sets[row.kind].add(row.code)
                    months = {k: v for k, v in fields.items() if len(k) == 10 and k[4] == "-"}
                    for key, value in months.items():
                        count["month_cells"] += 1
                        count["missing_month_cells"] += value is None
                        count["zero_month_cells"] += value == 0
                        count["negative_month_cells"] += value is not None and value < 0
                        identity = (row.code, key)
                        if identity in series[row.kind]:
                            supplier_stats[f"{row.kind}_duplicate_keys"] += 1
                        series[row.kind][identity] = value
                    total = fields.get("итого")
                    if row.kind == "sales_monthly" and total is not None:
                        count["totals_checked"] += 1
                        count["total_mismatches"] += abs(sum((v for v in months.values() if v is not None), Decimal(0)) - total) > Decimal("0.000001")
                    if row.kind == "transactions":
                        dt, qty = fields["дата"], fields["количество"]
                        first_date = min(first_date, dt) if first_date else dt
                        last_date = max(last_date, dt) if last_date else dt
                        count["missing_quantity"] += qty is None
                        if qty is not None:
                            count["positive_quantity"] += qty > 0
                            count["negative_quantity"] += qty < 0
                            count["zero_quantity"] += qty == 0
                            # A sign-inversion hypothesis, never an approved sales transformation.
                            transactions[(row.code, dt.strftime("%Y-%m-01"))] -= qty
                        transaction_units[row.code].add(fields.get("ед."))
                        count["warehouse:" + str(fields.get("склад"))] += 1
                        document = str(fields.get("документ", ""))
                        operation = document.split(str(fields.get("номер")))[0].strip()
                        count["operation:" + operation] += 1
                        if qty is not None:
                            sign = "positive" if qty > 0 else "negative" if qty < 0 else "zero"
                            count[f"operation_sign:{operation}:{sign}"] += 1
                    if row.kind == "constraints":
                        value = fields.get("кратность", fields.get("мин. разр. к отгр."))
                        constraints[row.code] = value
                        count["missing_constraint"] += value is None
                        count["nonpositive_constraint"] += value is not None and value <= 0
                    if row.kind == "sales_monthly" and "кратность" in fields:
                        sales_constraints[row.code] = fields["кратность"]
                    if row.kind == "summary":
                        physical, reserved, free = (fields.get(k) for k in ("остаток", "зарезервировано", "свободный остаток"))
                        if all(v is not None for v in (physical, reserved, free)):
                            count["free_stock_checked"] += 1
                            count["free_stock_mismatch"] += abs(physical - reserved - free) > Decimal("0.000001")
                        warehouses = [fields.get(k) for k in ("витрина", "остаток тз", "рц ект рыскулова", "розничный склад")]
                        if physical is not None and all(v is not None for v in warehouses):
                            count["warehouse_sum_checked"] += 1
                            count["warehouse_sum_mismatch"] += sum(warehouses) != physical
                        for key in ("кэф. роста", "кэф. сез-ти", "свободный остаток"):
                            value = fields.get(key)
                            count["negative:" + key] += value is not None and value < 0
            temporary.replace(target)
            item["date_min"] = first_date
            item["date_max"] = last_date
            item["kinds"] = {kind: {**count, "unique_codes": len(codes[kind])} for kind, count in counters.items()}
            # Scan every physical cell, including auxiliary sheets, formulas and Excel errors.
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                for sheet in book:
                    cell_counts = Counter()
                    fingerprint = hashlib.sha256()
                    for values in sheet.iter_rows():
                        populated = False
                        for cell in values:
                            if cell.value is None:
                                continue
                            populated = True
                            cell_counts["nonempty_cells"] += 1
                            cell_counts["formulas"] += cell.data_type == "f"
                            cell_counts["excel_errors"] += cell.data_type == "e"
                            fingerprint.update(dumps([cell.coordinate, cell.value]).encode())
                        cell_counts["nonempty_rows"] += populated
                    item["sheets"].append({"name": sheet.title, "max_row": sheet.max_row,
                                           "max_column": sheet.max_column, **cell_counts,
                                           "content_hash": fingerprint.hexdigest()})
            finally:
                book.close()
            if digest != file_hash(path):
                raise RuntimeError(f"Source changed during audit: {path.name}")
            report["files"].append(item)
            print(f"Profiled {supplier}: {path.name}", flush=True)
        comparisons = {}
        for left, right in (("sales_monthly", "constraints"), ("summary", "constraints"),
                            ("sales_monthly", "transactions"), ("summary", "sales_monthly")):
            comparisons[f"{left}__{right}"] = {
                "common": len(sets[left] & sets[right]), "left_only": len(sets[left] - sets[right]),
                "right_only": len(sets[right] - sets[left])}
        comparable = {k for k, v in series["sales_monthly"].items() if v is not None and k in transactions}
        matches = sum(abs(series["sales_monthly"][k] - transactions[k]) <= Decimal("0.000001") for k in comparable)
        direct_matches = sum(abs(series["sales_monthly"][k] + transactions[k]) <= Decimal("0.000001") for k in comparable)
        summary_common = {k for k, v in series["summary"].items() if v is not None and series["sales_monthly"].get(k) is not None}
        report["suppliers"][supplier] = {
            "code_coverage": {k: len(v) for k, v in sets.items()}, "joins": comparisons,
            "constraints_compared": len(constraints.keys() & sales_constraints.keys()),
            "constraint_conflicts": sum(constraints[k] != sales_constraints[k] for k in constraints.keys() & sales_constraints.keys()),
            "multi_unit_transaction_codes": sum(len(v) > 1 for v in transaction_units.values()),
            "monthly_vs_raw_transactions": {"compared": len(comparable), "equal": direct_matches,
                "different": len(comparable) - direct_matches,
                "note": "Raw signed quantities; operation types not filtered, scope not confirmed. Diagnostic only."},
            "monthly_vs_inverted_transactions": {"compared": len(comparable), "equal": matches,
                "different": len(comparable) - matches, "note": "Hypothesis only; scope and operation signs unconfirmed. Missing cells not treated as zero."},
            "summary_vs_monthly": {"compared": len(summary_common), "different": sum(
                abs(series["summary"][k] - series["sales_monthly"][k]) > Decimal("0.000001") for k in summary_common)},
            "duplicate_keys": dict(supplier_stats)}
    hashes = defaultdict(list)
    for item in report["files"]:
        for sheet in item["sheets"]:
            hashes[sheet["content_hash"]].append(f'{item["file"]} / {sheet["name"]}')
    report["identical_sheets"] = [v for v in hashes.values() if len(v) > 1]
    report["elapsed_seconds"] = round(time.perf_counter() - started, 2)
    (output / "profile.json").write_text(dumps(report) + "\n", encoding="utf-8")
    return report
