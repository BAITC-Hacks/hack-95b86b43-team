"""Read source rows without assigning unconfirmed business semantics."""
from dataclasses import dataclass
from datetime import date
from hashlib import sha256
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook

from ..normalization import identifier, label, month, number, timestamp

ADAPTER_VERSION = "1"
CODE_LABELS = {"код 1с", "номенклатура.код", "код"}


@dataclass(frozen=True)
class Layout:
    kind: str
    header_row: int
    headers: tuple
    code_column: int | None
    months: tuple[tuple[int, date], ...]


@dataclass(frozen=True)
class SourceRow:
    file: str
    sha256: str
    supplier: str
    sheet: str
    row: int
    kind: str
    code: str | None
    fields: dict
    errors: tuple[str, ...] = ()


def file_hash(path: Path) -> str:
    with path.open("rb") as stream:
        return sha256(stream.read()).hexdigest()


def detect_layout(rows: list[tuple]) -> Layout:
    for index, row in enumerate(rows):
        headers = tuple(label(v) for v in row)
        if "год" in headers and "янв" in headers:
            return Layout("seasonality", index + 1, tuple(row), None, ())
        code = next((i for i, value in enumerate(headers) if value in CODE_LABELS), None)
        if code is None:
            continue
        months = tuple((i, m) for i, value in enumerate(row) if (m := month(value)))
        if "дата" in headers and "количество" in headers:
            kind = "transactions"
        elif "свободный остаток" in headers:
            kind = "summary"
        elif any("поступление до" in h for h in headers):
            kind = "inbound"
        elif months:
            # Monthly inventory exports include a unit column, sales do not.
            kind = "inventory_monthly" if {"ед.", "ед.изм"} & set(headers) else "sales_monthly"
        elif {"кратность", "мин. разр. к отгр."} & set(headers):
            kind = "constraints"
        else:
            raise ValueError("Unrecognized product table")
        return Layout(kind, index + 1, tuple(row), code, months)
    raise ValueError("No recognized header in the first five rows")


def normalize_row(values: tuple, layout: Layout) -> tuple[str | None, dict, tuple[str, ...]]:
    fields, errors = {}, []
    code = identifier(values[layout.code_column]) if layout.code_column is not None else None
    numeric = {"количество", "кратность", "мин. разр. к отгр.", "остаток", "зарезервировано",
               "свободный остаток", "кэф. роста", "кэф. сез-ти", "витрина", "остаток тз",
               "рц ект рыскулова", "розничный склад", "итого"}
    for col, header in enumerate(layout.headers):
        key = label(header)
        if not key:
            continue
        value = values[col] if col < len(values) else None
        is_month = month(header)
        try:
            if is_month or key in numeric or "поступление до" in key or key.startswith("сэ в пути"):
                value = number(value)
            elif key == "дата":
                value = timestamp(value)
        except ValueError as exc:
            # Keep the raw invalid value alongside a blocking diagnostic.
            errors.append(f"column {col + 1}: {exc}")
        fields[is_month.isoformat() if is_month else key] = value
    return code, fields, tuple(errors)


def iter_workbook(path: Path, supplier: str) -> Iterator[SourceRow]:
    digest = file_hash(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            prefix = [next(rows, ()) for _ in range(5)]
            try:
                layout = detect_layout(prefix)
            except ValueError as exc:
                yield SourceRow(path.name, digest, supplier, sheet.title, 1, "unknown", None, {}, (str(exc),))
                continue
            from itertools import chain
            for row_number, values in enumerate(chain(prefix, rows), 1):
                if row_number <= layout.header_row or not any(v is not None for v in values):
                    continue
                if layout.kind == "seasonality":
                    # Preserve auxiliary rows; no automatic use as quantity seasonality.
                    yield SourceRow(path.name, digest, supplier, sheet.title, row_number,
                                    layout.kind, None, {str(i + 1): v for i, v in enumerate(values)})
                    continue
                raw_code = values[layout.code_column]
                if raw_code is None or label(raw_code) in {"итого", "всего"}:
                    # Headers/totals are audited separately, not turned into products.
                    yield SourceRow(path.name, digest, supplier, sheet.title, row_number,
                                    "non_product", None, {str(i + 1): v for i, v in enumerate(values)})
                    continue
                try:
                    code, fields, errors = normalize_row(values, layout)
                except ValueError as exc:
                    code, fields, errors = None, {}, (str(exc),)
                yield SourceRow(path.name, digest, supplier, sheet.title, row_number,
                                layout.kind, code, fields, errors)
    finally:
        book.close()
